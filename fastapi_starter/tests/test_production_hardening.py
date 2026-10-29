"""Regression tests for the production-readiness review (see the ALOS tracker,
'QA & Production Readiness' sheet). Each test names the bug it pins down."""
import io
import time
import types
import zipfile

import pytest
from conftest import wait_for_chapter_ready


# --------------------------------------------------------------------------- #
# Video content: the script must TEACH, not give stage directions
# --------------------------------------------------------------------------- #
DIRECTOR_PHRASES = [
    "open with a relatable", "on-screen diagram", "illustrate ", "leave students with",
    "record/animate", "short visual example", "so students see why",
]

ENGLISH_CHAPTER = (
    "Photosynthesis is the process by which green plants make their own food using sunlight. "
    "Photosynthesis occurs mainly in the leaves of plants. During photosynthesis, plants take in "
    "carbon dioxide from the air. Water is absorbed from the soil through the roots. "
    "Chlorophyll is a green pigment in the leaves. Chlorophyll absorbs sunlight to power photosynthesis."
)
HINDI_CHAPTER = (
    "पौधे सूर्य के प्रकाश से अपना भोजन बनाते हैं। इस प्रक्रिया को प्रकाश संश्लेषण कहते हैं। "
    "पत्तियों में क्लोरोफिल होता है। क्लोरोफिल सूर्य का प्रकाश सोखता है। "
    "प्रकाश संश्लेषण में पौधे कार्बन डाइऑक्साइड और जल का उपयोग करते हैं।"
)


def _chapter(name="Photosynthesis", subject="Science", cls="Class 7"):
    return types.SimpleNamespace(chapter_name=name, subject=subject, class_name=cls)


def test_video_script_teaches_chapter_content_not_stage_directions():
    """Bug: scenes were instructions to the video maker ('Open with a relatable
    question...', 'using a simple on-screen diagram') which were then printed on
    the slides AND read aloud to students."""
    from app.services.ai_service import analyze_chapter_text
    from app.services.content_generation_service import generate_content

    analysis = analyze_chapter_text(ENGLISH_CHAPTER)
    body = generate_content(_chapter(), "video", analysis["key_concepts"], analysis)["body"].lower()

    for phrase in DIRECTOR_PHRASES:
        assert phrase not in body, f"script still contains a stage direction: {phrase!r}"
    # ...and it contains the chapter's own explanatory sentence
    assert "photosynthesis is the process by which green plants make their own food" in body
    assert "scene 1 (introduction)" in body and "(recap)" in body


def test_video_script_fallback_without_analysis_is_still_not_a_stage_direction():
    from app.services.content_generation_service import generate_content
    body = generate_content(_chapter(), "video", ["Chlorophyll", "Sunlight"])["body"].lower()
    for phrase in DIRECTOR_PHRASES:
        assert phrase not in body


def test_hindi_chapter_is_analyzed_and_gets_a_hindi_script():
    """Bug: WORD_RE only matched A-Z, so a Hindi chapter failed with 'no
    meaningful words' and never reached the video step."""
    from app.services.ai_service import analyze_chapter_text
    from app.services.content_generation_service import generate_content

    analysis = analyze_chapter_text(HINDI_CHAPTER)
    assert analysis["key_concepts"], "Hindi concepts should be extracted"
    assert "के" not in analysis["key_concepts"] and "है" not in analysis["key_concepts"]  # stopwords removed
    assert analysis["concept_sentences"], "Hindi sentences should be matched to concepts (danda split)"

    body = generate_content(_chapter("प्रकाश संश्लेषण", "Hindi"), "video", analysis["key_concepts"], analysis)["body"]
    assert "आज हम" in body and "दोहराएँ" in body   # Hindi narration templates, not English


def test_key_concepts_skip_generic_verbs():
    from app.services.ai_service import analyze_chapter_text
    concepts = [c.lower() for c in analyze_chapter_text(ENGLISH_CHAPTER)["key_concepts"]]
    assert "produced" not in concepts and "occurs" not in concepts


# --------------------------------------------------------------------------- #
# Matching: subject alone must never qualify an item
# --------------------------------------------------------------------------- #
def _library_item(db, models, concepts, content_type="video", subject="Science", cls="Class 7", media="/media/videos/x.mp4"):
    tags = []
    for name in concepts + [subject.lower(), cls.lower()]:
        t = db.query(models.ContentTag).filter_by(name=name).first() or models.ContentTag(name=name)
        db.add(t)
        tags.append(t)
    item = models.ContentItem(
        title="Photosynthesis - Explainer", content_type=content_type, subject=subject, class_name=cls,
        body="x" * 60, media_url=media, is_generated=True, is_seed_placeholder=False, tags=tags,
    )
    db.add(item)
    db.commit()
    return item


def test_chapter_with_different_concepts_does_not_get_another_chapters_video():
    """Bug: MIN_MATCH_SCORE (0.25) < SUBJECT_WEIGHT (0.3), so a same-subject
    item matched with ZERO concept overlap: a Digestive System chapter was handed
    the Photosynthesis video, and (generated videos are saved to the library)
    every later chapter of that subject got it too."""
    from app import models
    from app.database import SessionLocal
    from app.services.matching_service import find_best_matches, evaluate_matches

    db = SessionLocal()
    try:
        unique = f"zq{int(time.time() * 1000)}"
        item = _library_item(db, models, [f"{unique}chlorophyll", f"{unique}sunlight", f"{unique}photosynthesis"])

        other = _chapter("Digestive System", "Science", "Class 7")
        res = evaluate_matches(find_best_matches(db, other, ["Stomach", "Intestine", "Enzymes", "Saliva"]))
        assert res["video"]["item"] is None or res["video"]["item"].id != item.id

        same = _chapter("Photosynthesis", "Science", "Class 7")
        res = evaluate_matches(find_best_matches(
            db, same, [f"{unique}Chlorophyll", f"{unique}Sunlight", "Something Else"]))
        assert res["video"]["item"] is not None and res["video"]["item"].id == item.id
        assert res["video"]["is_valid"]
    finally:
        db.close()


def test_one_shared_concept_is_not_enough_when_chapter_has_many():
    from app import models
    from app.database import SessionLocal
    from app.services.matching_service import find_best_matches

    db = SessionLocal()
    try:
        unique = f"yq{int(time.time() * 1000)}"
        item = _library_item(db, models, [f"{unique}plants"], subject="Science", cls="Class 9")
        ch = _chapter("Other", "Science", "Class 9")
        res = find_best_matches(db, ch, [f"{unique}Plants", "Gravity", "Orbits", "Tides"])
        assert res["video"]["item"] is None or res["video"]["item"].id != item.id
    finally:
        db.close()


def test_tag_matching_is_word_based_not_substring():
    """Bug: substring matching made the tag 'air' match the concept 'chair'."""
    from app.services.matching_service import _concept_matches_tag
    assert not _concept_matches_tag("chair", "air")
    assert not _concept_matches_tag("excellent", "cell")
    assert _concept_matches_tag("Water Cycle", "cycle")
    assert _concept_matches_tag("Plants", "plant")


def test_seed_cartoon_placeholder_is_not_a_valid_match():
    """Bug: seeded cartoons point at a fake CDN but only seeded *videos* were
    rejected, so chapters got broken cartoon image links."""
    from app import models
    from app.services.matching_service import validate_match

    tag = models.ContentTag(name="photosynthesis")
    fake = models.ContentItem(
        title="x", content_type="cartoon", subject="Science", class_name="Class 6",
        media_url="https://cdn.alos.app/cartoons/photosynthesis.png", is_seed_placeholder=True, tags=[tag],
    )
    ok, reason = validate_match(fake, 1.0)
    assert ok is False and "placeholder" in reason


# --------------------------------------------------------------------------- #
# Video serving
# --------------------------------------------------------------------------- #
def test_media_supports_range_requests():
    """Bug: Starlette 0.38's StaticFiles ignored Range headers: iOS/Safari
    refuse to play mp4 without them and nothing could seek."""
    import os
    from fastapi.testclient import TestClient
    from app.main import app
    from app.config import settings

    d = os.path.join(settings.GENERATED_MEDIA_DIR, "videos")
    os.makedirs(d, exist_ok=True)
    with open(os.path.join(d, "range_test.bin"), "wb") as f:
        f.write(bytes(range(256)) * 40)

    with TestClient(app) as c:
        r = c.get("/media/videos/range_test.bin", headers={"Range": "bytes=10-19"})
    assert r.status_code == 206
    assert r.headers["content-range"].startswith("bytes 10-19/")
    assert len(r.content) == 10


def test_rendered_video_has_captions_poster_and_faststart(tmp_path):
    from app.services import video_rendering_service as v
    script = (
        "VIDEO SCRIPT - T\n\n"
        "Scene 1 (Introduction): Today we learn about plants.\n"
        "Scene 2 (Key idea: Leaves): Leaves make food from sunlight.\n"
    )
    rel = v.render_video("Plants", script, str(tmp_path))
    mp4 = tmp_path / rel
    assert mp4.exists() and mp4.read_bytes()[4:8] == b"ftyp"
    vtt = mp4.with_suffix(".vtt").read_text(encoding="utf-8")
    assert vtt.startswith("WEBVTT") and "Leaves make food from sunlight." in vtt
    assert mp4.with_suffix(".jpg").exists()
    # faststart => the 'moov' atom comes before 'mdat' so playback can begin immediately
    data = mp4.read_bytes()
    assert 0 < data.find(b"moov") < data.find(b"mdat")

    v.delete_rendered(rel, str(tmp_path))
    assert not mp4.exists() and not mp4.with_suffix(".vtt").exists() and not mp4.with_suffix(".jpg").exists()
    assert not list((tmp_path / "videos").glob("_tmp_*")), "temp work dir must be cleaned up"


def test_failed_render_leaves_no_files_behind(tmp_path, monkeypatch):
    from app.services import video_rendering_service as v

    def boom(*a, **k):
        raise v.VideoRenderingError("ffmpeg failed: simulated")
    monkeypatch.setattr(v, "_run_ffmpeg", boom)
    with pytest.raises(v.VideoRenderingError):
        v.render_video("X", "Scene 1 (Introduction): hello there", str(tmp_path))
    leftovers = [p for p in (tmp_path / "videos").iterdir()] if (tmp_path / "videos").exists() else []
    assert leftovers == []


# --------------------------------------------------------------------------- #
# Pipeline robustness
# --------------------------------------------------------------------------- #
def test_other_writes_are_not_blocked_while_a_video_renders(client, make_user, tmp_path, monkeypatch):
    """Bug: the whole pipeline ran in ONE transaction, so on SQLite the write
    lock was held for the entire render and every other write hit 'database is
    locked' once a render outlasted the 15 s timeout."""
    from docx import Document
    from app.routers import chapters_router
    from app import models
    from app.database import SessionLocal

    outcome = {}

    def fake_render(chapter_name, body, out_dir):
        db2 = SessionLocal()
        started = time.monotonic()
        try:
            db2.add(models.ErrorLog(path="/probe", method="GET", error_type="Probe", error_message="write during render", traceback="-"))
            db2.commit()
            outcome["ok"], outcome["secs"] = True, time.monotonic() - started
        except Exception as exc:                      # noqa: BLE001
            outcome["ok"], outcome["err"] = False, repr(exc)
        finally:
            db2.close()
        return "videos/fake_probe.mp4"

    monkeypatch.setattr(chapters_router, "render_video", fake_render)

    doc = Document()
    doc.add_paragraph("Zorbulation theory explains how quibbling frobnicators synchronise. "
                      "Quibbling frobnicators need zorbulation to synchronise their wibbles. "
                      "Wibbles drift when frobnicators stop quibbling.")
    path = tmp_path / "c.docx"
    doc.save(str(path))

    _, headers = make_user("teacher")
    with open(path, "rb") as f:
        r = client.post("/chapters/upload", headers=headers,
                        data={"class_name": "Class 11", "subject": "Fictional Physics", "chapter_name": "Zorbulation"},
                        files={"file": ("c.docx", f, "application/octet-stream")})
    assert r.status_code == 201
    assert wait_for_chapter_ready(client, headers, r.json()["id"], timeout=60) == "ready"
    assert outcome.get("ok") is True, f"a write was blocked/failed during the render: {outcome}"
    assert outcome["secs"] < 3


def test_failed_db_write_deletes_the_rendered_video(client, make_user, tmp_path, monkeypatch):
    """Bug: if the pipeline failed after rendering, the .mp4 stayed on disk forever."""
    from docx import Document
    from app.routers import chapters_router

    deleted = []
    monkeypatch.setattr(chapters_router, "render_video", lambda *a, **k: "videos/orphan_probe.mp4")
    monkeypatch.setattr(chapters_router, "delete_rendered", lambda rel, d: deleted.append(rel))

    def boom(*a, **k):
        raise RuntimeError("simulated DB failure while saving")
    monkeypatch.setattr(chapters_router, "_persist_matching", boom)

    doc = Document()
    doc.add_paragraph("Glimmerfold dynamics describe how snorkling wobblers cascade. "
                      "Snorkling wobblers cascade through glimmerfold layers. Glimmerfold layers hum.")
    path = tmp_path / "g.docx"
    doc.save(str(path))
    _, headers = make_user("teacher")
    with open(path, "rb") as f:
        r = client.post("/chapters/upload", headers=headers,
                        data={"class_name": "Class 11", "subject": "Fictional Physics", "chapter_name": "Glimmerfold"},
                        files={"file": ("g.docx", f, "application/octet-stream")})
    assert wait_for_chapter_ready(client, headers, r.json()["id"], timeout=60) == "failed"
    assert deleted == ["videos/orphan_probe.mp4"]


def test_upload_rejects_file_whose_content_does_not_match_extension(client, make_user):
    """Bug: only the extension was checked, so any bytes named chapter.pdf reached the parsers."""
    _, headers = make_user("teacher")
    r = client.post("/chapters/upload", headers=headers,
                    data={"class_name": "Class 6", "subject": "Science", "chapter_name": "Fake"},
                    files={"file": ("chapter.pdf", io.BytesIO(b"MZ\x90\x00 this is not a pdf"), "application/pdf")})
    assert r.status_code == 422 and "does not look like" in r.json()["detail"]

    r = client.post("/chapters/upload", headers=headers,
                    data={"class_name": "Class 6", "subject": "Science", "chapter_name": "Fake2"},
                    files={"file": ("chapter.docx", io.BytesIO(b"plain text, not a zip"), "application/octet-stream")})
    assert r.status_code == 422

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "a zip that is not a docx")
    r = client.post("/chapters/upload", headers=headers,
                    data={"class_name": "Class 6", "subject": "Science", "chapter_name": "Fake3"},
                    files={"file": ("chapter.docx", io.BytesIO(buf.getvalue()), "application/octet-stream")})
    assert r.status_code == 422


def test_failed_chapter_can_be_reprocessed(client, make_user, sample_docx):
    """Bug: a failed chapter was a dead end - the only option was uploading again."""
    from app import models
    from app.database import SessionLocal

    _, headers = make_user("teacher")
    with open(sample_docx, "rb") as f:
        r = client.post("/chapters/upload", headers=headers,
                        data={"class_name": "Class 6", "subject": "Science", "chapter_name": "Reprocess Me"},
                        files={"file": ("chapter.docx", f, "application/octet-stream")})
    cid = r.json()["id"]
    assert wait_for_chapter_ready(client, headers, cid, timeout=90) == "ready"

    # a 'ready' chapter can't be reprocessed
    assert client.post(f"/chapters/{cid}/reprocess", headers=headers).status_code == 409

    db = SessionLocal()
    ch = db.query(models.Chapter).get(cid)
    ch.status, ch.status_message = "failed", "simulated failure"
    db.commit()
    db.close()

    # someone else can't reprocess it
    _, other = make_user("teacher")
    assert client.post(f"/chapters/{cid}/reprocess", headers=other).status_code == 403

    r = client.post(f"/chapters/{cid}/reprocess", headers=headers)
    assert r.status_code == 202
    assert wait_for_chapter_ready(client, headers, cid, timeout=90) == "ready"


def test_stuck_chapters_are_requeued_on_startup(client, make_user, sample_docx, monkeypatch):
    """Bug: background jobs live in process memory - a restart mid-render left
    chapters 'processing' forever."""
    import datetime
    from app import models
    from app.database import SessionLocal
    from app.routers import chapters_router

    _, headers = make_user("teacher")
    with open(sample_docx, "rb") as f:
        r = client.post("/chapters/upload", headers=headers,
                        data={"class_name": "Class 6", "subject": "Science", "chapter_name": "Stuck Chapter"},
                        files={"file": ("chapter.docx", f, "application/octet-stream")})
    cid = r.json()["id"]
    wait_for_chapter_ready(client, headers, cid, timeout=90)

    db = SessionLocal()
    ch = db.query(models.Chapter).get(cid)
    ch.status = "processing"
    db.commit()
    # `updated_at` is bumped by onupdate - set it back explicitly in a second statement
    db.query(models.Chapter).filter(models.Chapter.id == cid).update(
        {"updated_at": datetime.datetime.utcnow() - datetime.timedelta(hours=2)}, synchronize_session=False)
    db.commit()
    db.close()

    processed = []
    monkeypatch.setattr(chapters_router, "_process_chapter", lambda chapter_id: processed.append(chapter_id))
    queued = chapters_router.recover_stuck_chapters()
    assert queued >= 1
    deadline = time.monotonic() + 5
    while cid not in processed and time.monotonic() < deadline:
        time.sleep(0.1)
    assert cid in processed


def test_chapter_list_is_paginated(client, make_user):
    _, headers = make_user("teacher")
    assert client.get("/chapters/?limit=1&offset=0", headers=headers).status_code == 200
    assert client.get("/chapters/?limit=0", headers=headers).status_code == 422
    assert client.get("/chapters/?limit=1000", headers=headers).status_code == 422


# --------------------------------------------------------------------------- #
# Operations
# --------------------------------------------------------------------------- #
def test_every_response_carries_a_request_id(client):
    r = client.get("/health")
    assert r.headers.get("x-request-id")
    r = client.get("/health", headers={"X-Request-ID": "trace-abc_123"})
    assert r.headers["x-request-id"] == "trace-abc_123"
    r = client.get("/health", headers={"X-Request-ID": "bad id with spaces <script>"})
    assert r.headers["x-request-id"] != "bad id with spaces <script>"


def test_health_reports_video_capabilities(client):
    body = client.get("/health").json()
    assert body["status"] == "ok"
    assert set(body["video"]) == {"ffmpeg", "narration"}


def test_admin_signup_defaults_to_closed_in_production(monkeypatch):
    """Bug: render.yaml deployed with ALLOW_ADMIN_SIGNUP unset -> default true -> anyone could register as admin."""
    import importlib
    from app import config
    monkeypatch.setenv("ENV", "production")
    monkeypatch.setenv("SECRET_KEY", "x" * 48)
    monkeypatch.delenv("ALLOW_ADMIN_SIGNUP", raising=False)
    try:
        reloaded = importlib.reload(config)
        assert reloaded.settings.ALLOW_ADMIN_SIGNUP is False
        assert reloaded.settings.AUTO_CREATE_TABLES is False
        monkeypatch.setenv("ALLOW_ADMIN_SIGNUP", "true")
        assert importlib.reload(config).settings.ALLOW_ADMIN_SIGNUP is True   # explicit opt-in still works
    finally:
        monkeypatch.undo()
        importlib.reload(config)


def test_jwt_without_expiry_or_with_alg_none_is_rejected(client, make_user):
    """Tokens must carry exp + sub and use the configured algorithm; python-jose was
    replaced by PyJWT (unmaintained library, unfixed JWT CVE)."""
    import datetime
    import jwt
    from app.config import settings

    user, headers = make_user("teacher")
    assert client.get("/auth/me", headers=headers).status_code == 200   # sanity: a real token works

    no_exp = jwt.encode({"sub": str(user["id"]), "type": "access"}, settings.SECRET_KEY, algorithm="HS256")
    assert client.get("/auth/me", headers={"Authorization": f"Bearer {no_exp}"}).status_code == 401

    exp = datetime.datetime.utcnow() + datetime.timedelta(minutes=5)
    alg_none = jwt.encode({"sub": str(user["id"]), "type": "access", "exp": exp}, None, algorithm="none")
    assert client.get("/auth/me", headers={"Authorization": f"Bearer {alg_none}"}).status_code == 401

    wrong_key = jwt.encode({"sub": str(user["id"]), "type": "access", "exp": exp}, "x" * 40, algorithm="HS256")
    assert client.get("/auth/me", headers={"Authorization": f"Bearer {wrong_key}"}).status_code == 401
